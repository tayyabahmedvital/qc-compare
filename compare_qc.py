"""Compare a results-flagger table against the ground-truth PASS/FAIL table.

    python compare_qc.py --truth inputs/<truth>.csv --flagger inputs/<flagger>.csv --check absqc-uf-only

Writes three files per check into --out (default `outputs/`):
    fail comparison - <check>.csv                      every run
    fail comparison - <check> (CAUGHT-MISS only).csv   runs with at least one CAUGHT or MISS
    fail comparison - <check>.xlsx                     both pages, coloured

Cell labels (truth vs flagger):
    CAUGHT     truth FAIL, flagger FAIL        MISS   truth FAIL, flagger PASS
    GOOD       truth not FAIL, flagger PASS    WRONG  truth not FAIL, flagger FAIL (false positive)
    Not Found  the flagger has no value for that run/assay, or the truth has no row for the run
Rows follow the flagger's run order; a blank truth cell counts as "not FAIL".
"""

import argparse
import csv
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill

# The uF disc layout: assay -> structure. Truth columns for any other assay (HT, IA) are ignored.
STRUCTURE = {
    "HDL": "A", "ALP": "A", "CREA": "A",
    "GGT": "B", "ALT": "B", "TBIL": "B", "AST": "B",
    "HBA1C1": "C", "HBA1C2": "C",
    "CHOL": "D (1R)", "ALB": "D (1R)", "TRIG": "D (1R)", "GLU": "D (1R)", "CA": "D (1R)",
    "TP": "D (2R)", "BUN": "D (2R)", "TCO2": "D (2R)",
    "CL": "E", "NA": "E", "K": "E",
}
LABELS = ["CAUGHT", "MISS", "GOOD", "WRONG", "Not Found"]
DIV0 = "#DIV/0!"
STATS_ROWS = 7  # the stats block fills rows 1-7 of the sheet
FIRST_RUN_ROW = STATS_ROWS + 3  # after the struct and run_code header rows
FILLS = {  # label -> (font colour, fill colour)
    "MISS": ("000000", "FFE599"), "WRONG": ("000000", "F4B6B6"),
    "GOOD": ("999999", None), "Not Found": ("999999", "D9D9D9"),
}


def read_grid(path):
    return pd.read_csv(path, header=None, dtype=str, keep_default_na=False)


def load_truth(path):
    """Truth CSV (any export layout) -> PASS/FAIL frame, run code x assay, uF assays only.

    The assay-name row is whichever of the first rows holds the most known assays; the
    layouts differ only in how many header rows sit above it.
    """
    raw = read_grid(path)
    header = max(range(6), key=lambda i: raw.iloc[i].isin(STRUCTURE).sum())
    body = raw.iloc[header + 1:]
    truth = body[body[0] != ""].set_index(0)
    truth.columns = raw.iloc[header, 1:].tolist()
    if not truth.index.is_unique:
        raise ValueError(f"duplicate run codes in {path}: {truth.index[truth.index.duplicated()].tolist()}")
    return truth.loc[:, truth.columns.isin(STRUCTURE)]


def load_flagger(path):
    """Flagger CSV (`run_code` plus one PASS/FAIL column per assay) -> frame, run x assay."""
    flagger = pd.read_csv(path, dtype=str, keep_default_na=False, index_col="run_code")
    unknown = set(np.unique(flagger.to_numpy())) - {"PASS", "FAIL", ""}
    if unknown:
        raise ValueError(f"{path}: expected PASS/FAIL, found {sorted(unknown)}")
    return flagger


def compare(truth, flagger, pad=0):
    """Label every (run, assay) cell. `pad` appends blank-header Not Found columns, as in the sheets."""
    t = truth.reindex(flagger.index).fillna("")
    f = flagger.reindex(columns=truth.columns).fillna("")
    not_found = (f == "") | ~flagger.index.isin(truth.index)[:, None]
    failed = t == "FAIL"
    flagged = f == "FAIL"
    cells = np.select([not_found, failed & flagged, failed, flagged],
                      ["Not Found", "CAUGHT", "MISS", "WRONG"], default="GOOD")
    assays = truth.columns.tolist() + [""] * pad
    cells = np.hstack([cells, np.full((len(cells), pad), "Not Found")])
    columns = pd.MultiIndex.from_arrays([[STRUCTURE.get(a, "") for a in assays], assays])
    return pd.DataFrame(cells, index=flagger.index, columns=columns)


def ratio(num, den):
    return num / den if den else DIV0


def stats(labels, runs_label):
    """The seven stats rows, one value per column of `labels`."""
    n = len(labels)
    c = pd.DataFrame({label: (labels == label).sum().to_numpy() for label in LABELS})
    failures = c["CAUGHT"] + c["MISS"]
    return [
        [runs_label, *[n] * len(c)],
        ["Failure count", *failures],
        ["Detected failures (correct)", *c["CAUGHT"]],
        ["False positives", *c["WRONG"]],
        ["Prevalence rate", *failures / n],
        ["Sensitivity - True positive detection rate", *map(ratio, c["CAUGHT"], failures)],
        ["Specificity", *map(ratio, c["GOOD"], c["GOOD"] + c["WRONG"])],
    ]


def page(stat_rows, labels):
    """Stats block, struct/assay header rows, then one row per run: the layout of the sheet."""
    header = [["struct", *labels.columns.get_level_values(0)],
              ["run_code", *labels.columns.get_level_values(1)]]
    return stat_rows + header + [[run, *row] for run, row in zip(labels.index, labels.to_numpy())]


def write_csv(rows, path):
    def text(v):
        return f"{v:.2%}" if isinstance(v, float) else v

    with open(path, "w", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows([text(v) for v in row] for row in rows)


def add_sheet(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=FIRST_RUN_ROW, min_col=2):
        for cell in row:
            font, fill = FILLS.get(cell.value, ("000000", None))
            cell.font = Font(color=font, bold=cell.value == "CAUGHT")
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    for row in sheet.iter_rows(min_row=5, max_row=7, min_col=2):  # prevalence, sensitivity, specificity
        for cell in row:
            cell.number_format = "0.00%"
    sheet.conditional_formatting.add(
        f"B6:{sheet.cell(7, sheet.max_column).coordinate}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F4B6B6",
                       end_type="num", end_value=1, end_color="93C47D"))
    sheet.freeze_panes = f"B{FIRST_RUN_ROW}"
    sheet.column_dimensions["A"].width = 42


def write_all(labels, check, out_dir, runs_label="Total runs with data"):
    """Write both pages as CSVs and as sheets of one workbook; returns the run count per page."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stat_rows = stats(labels, runs_label)
    needs_review = labels[labels.isin(["CAUGHT", "MISS"]).any(axis=1)]
    pages = [("all runs", f"fail comparison - {check}", labels),
             ("CAUGHT-MISS only", f"fail comparison - {check} (CAUGHT-MISS only)", needs_review)]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_title, file_stem, page_labels in pages:
        rows = page(stat_rows, page_labels)
        write_csv(rows, out_dir / f"{file_stem}.csv")
        add_sheet(workbook, sheet_title, rows)
    workbook.save(out_dir / f"fail comparison - {check}.xlsx")
    return {file_stem: len(page_labels) for _, file_stem, page_labels in pages}


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--truth", required=True, help="truth PASS/FAIL CSV (well_response, truth (uf only), ...)")
    parser.add_argument("--flagger", required=True, help="flagger CSV: run_code + one PASS/FAIL column per assay")
    parser.add_argument("--check", required=True, help="name used in the output files, e.g. absqc-uf-only")
    parser.add_argument("--out", default="outputs", help="output folder (default: outputs)")
    parser.add_argument("--runs-label", default="Total runs with data", help="label of the first stats row")
    parser.add_argument("--pad", type=int, default=4,
                        help="trailing blank-header columns the sheets carry (default 4; 0 to drop)")
    args = parser.parse_args()

    labels = compare(load_truth(args.truth), load_flagger(args.flagger), args.pad)
    for name, n_runs in write_all(labels, args.check, args.out, args.runs_label).items():
        print(f"{n_runs:>4} runs  {name}")


if __name__ == "__main__":
    main()
